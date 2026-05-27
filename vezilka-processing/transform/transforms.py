import os
import fitz
import uuid
import requests
import tempfile
import pandas as pd
import textract
import logging

from image_to_text.image_to_text import load_image_and_transform
from pptx import Presentation
from docx import Document
from transform.transform_utils import *
from datasets import Dataset

def transform_pdf(file_path, url, row):
    doc = fitz.open(file_path)

    for page_num, page in enumerate(doc):
        text = page.get_text("text").strip()

        if not text and OCR_PAGES:
            text = ocr_page(page) # the pdf might be fully scanned so get text from page with ocr

        chunks = chunk_text(text)

        for i, chunk in enumerate(chunks):
            yield DatasetRow(
                id=str(uuid.uuid4()),
                text=chunk,
                source=url,
                page=page_num,
                chunk=i,
                topic=row.topic,
                description=row.description,
                file_type=row.type,
                dialect=row.dialect
            )

def transform_docx_text(docx_path):
    doc = Document(docx_path)

    paragraphs = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)

def transform_doc_text(doc_path):
    text = textract.process(doc_path).decode("utf-8")

    return text

def transform_document(file_path, url, row):
    extension = os.path.splitext(url)[1].lower()

    try:
        if extension == ".docx":
            text = transform_docx_text(file_path)

        elif extension == ".doc":
            text = transform_doc_text(file_path)

        else:
            logging.warning(f"Skipping unsupported file type: {extension} | {url}")
            return

        chunks = chunk_text(text)

        for i, chunk in enumerate(chunks):
            yield DatasetRow(
                id=str(uuid.uuid4()),
                text=chunk,
                source=url,
                chunk=i,
                topic=row.topic,
                description=row.description,
                file_type=row.type,
                dialect=row.dialect
            )

    except Exception as e:
        logging.error(f"Failed processing {url}: {e}")
        return

def transform_txt(file_path, url, row):
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()

    text = text.strip()

    chunks = chunk_text(text)

    for i, chunk in enumerate(chunks):
        yield DatasetRow(
            id=str(uuid.uuid4()),
            text=chunk,
            source=url,
            chunk=i,
            topic=row.topic,
            description=row.description,
            file_type=row.type,
            dialect=row.dialect
        )

def transform_pptx(file_path, url, row):
    presentation = Presentation(file_path)

    for slide_num, slide in enumerate(presentation.slides):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        full_text = "\n".join(slide_text).strip()

        if not full_text:
            continue

        chunks = chunk_text(full_text)

        for i, chunk in enumerate(chunks):
            yield DatasetRow(
                id=str(uuid.uuid4()),
                text=chunk,
                source=url,
                chunk=i,
                topic=row.topic,
                description=row.description,
                file_type=row.type,
                dialect=row.dialect
            )

def transform_image(file_path, url, row):
    text = load_image_and_transform(file_path)

    chunks = chunk_text(text)

    for i, chunk in enumerate(chunks):
        yield DatasetRow(
            id=str(uuid.uuid4()),
            text=chunk,
            source=url,
            chunk=i,
            topic=row.topic,
            description=row.description,
            file_type=row.type,
            dialect=row.dialect
        )

def transform_video_audio(file_path, url, row):
    text = row.text

    chunks = chunk_text(text)

    for i, chunk in enumerate(chunks):
        yield DatasetRow(
            id=str(uuid.uuid4()),
            text=chunk,
            source=url,
            chunk=i,
            topic=row.topic,
            description=row.description,
            file_type=row.type,
            dialect=row.dialect
        )

def download_file(url):
    r = requests.get(url)
    r.raise_for_status()

    suffix = os.path.splitext(url)[1]

    tmp = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=suffix
    )

    tmp.write(r.content)
    tmp.close()

    return tmp.name

def process_file(file_row):
    file_url = (
        f"{os.getenv('VITE_BACKEND_URL').rstrip('/')}/"
        f"{file_row.file_url.lstrip('/')}"
    )

    ext = os.path.splitext(file_url)[1].lower()

    handler = SUPPORTED_EXTENSIONS.get(ext)

    if handler is None:
        logging.warning(
            f"Unsupported file type: {ext} | {file_url}"
        )
        return

    try:
        file_path = download_file(file_url)

        yield from handler(file_path, file_url, file_row)

    except Exception as e:
        logging.exception(
            f"Failed processing file: {file_url} | {e}"
        )

SUPPORTED_EXTENSIONS = {
    ".pdf": transform_pdf,

    ".doc": transform_document,
    ".docx": transform_document,

    ".txt": transform_txt,

    ".pptx": transform_pptx,

    ".png": transform_image,
    ".jpg": transform_image,
    ".jpeg": transform_image,

    ".mp3": transform_video_audio,
    ".wav": transform_video_audio,
    ".ogg": transform_video_audio,
    ".flac": transform_video_audio,
    ".aac": transform_video_audio,
    ".m4a": transform_video_audio,

    ".mp4": transform_video_audio,
    ".mov": transform_video_audio,
    ".avi": transform_video_audio,
    ".mkv": transform_video_audio,
    ".webm": transform_video_audio,
    ".wmv": transform_video_audio,
}